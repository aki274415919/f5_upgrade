from pathlib import Path
import re
try:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
except Exception:  # pragma: no cover - allow running without openpyxl
    Workbook = None
    PatternFill = None

from .messages import MESSAGES

ERROR_KEYS = [
    "expired", "invalid", "fail", "error", "critical", "alarm", "unavailable",
    "inactive", "offline", "lost", "unreachable", "down", "degraded",
]
WARN_KEYS = ["warn", "deprecated", "high", "exceeded", "mismatch"]


def analyse(cmd: str, text: str):
    """Determine severity and message based on command output."""
    lower = text.lower()
    severity = "normal"
    msg = MESSAGES["default_ok"]
    # disk usage check
    if "show sys disk" in cmd:
        for num in re.findall(r"(\d+)%", text):
            if int(num) >= 80:
                severity = "warning"
                msg = MESSAGES["disk_usage"]
                break
    if "deprecated" in cmd and text.strip():
        severity = "warning"
        msg = MESSAGES["default_warning"]
    if any(k in lower for k in ERROR_KEYS):
        severity = "critical"
        msg = MESSAGES["default_warning"]
    elif any(k in lower for k in WARN_KEYS):
        severity = "warning"
        msg = MESSAGES["default_warning"]
    return severity, msg


def apply_highlight(text, rules):
    def repl(match):
        word = match.group(0)
        color = rules[word.lower()]
        return f"<mark style='background-color:{color}'>{word}</mark>"

    pattern = re.compile("|".join(re.escape(k) for k in rules), re.IGNORECASE)
    return pattern.sub(lambda m: repl(m), text)


def save_html(results, path):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write("<html><meta charset='utf-8'><body>\n")
        f.write("<table border='1' cellspacing='0' cellpadding='5'>\n")
        f.write(
            "<tr><th>序号/No</th><th>检查项/Item</th><th>设备IP/IP"  # header part 1
            "</th><th>输出/Output</th><th>说明/Notes</th></tr>\n"
        )
        for idx, item, ip, out, _ in results:
            sev, msg = analyse(item, out)
            rules = {
                k: ("#ff6b6b" if k in ERROR_KEYS else "#ffd93b")
                for k in ERROR_KEYS + WARN_KEYS
                if k in out.lower()
            }
            out_html = apply_highlight(out, rules) if rules else out
            note = "<br/>".join([msg[lang] for lang in ("zh", "ja", "en")])
            f.write(
                f"<tr><td>{idx}</td><td>{item}</td><td>{ip}</td>"
                f"<td>{out_html}</td><td>{note}</td></tr>\n"
            )
        f.write("</table></body></html>")


def save_excel(results, path):
    if Workbook is None:
        import csv

        with open(path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["序号", "检查项", "设备IP", "输出", "需要关注的内容（日文）"])
            for idx, item, ip, out, _ in results:
                _, msg = analyse(item, out)
                writer.writerow([idx, item, ip, out, msg["ja"]])
        return

    wb = Workbook()
    ws = wb.active
    ws.append(["序号", "检查项", "设备IP", "输出", "需要关注的内容（日文）"])
    colors = {
        "critical": "FFC7CE",  # red
        "warning": "FFEB9C",  # yellow
        "normal": "C6EFCE",  # green
    }
    for idx, item, ip, out, _ in results:
        sev, msg = analyse(item, out)
        ws.append([idx, item, ip, out, msg["ja"]])
        fill = PatternFill(start_color=colors[sev], end_color=colors[sev], fill_type="solid")
        for cell in ws[ws.max_row]:
            cell.fill = fill
    wb.save(path)
